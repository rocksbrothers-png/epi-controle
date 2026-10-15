"""Registry de fontes do Centro de Migração (ADR-0003 §2.2).

Toda fonte — XLSX, CSV, ODS, JSON, XML, TXT e (fases 5/6) banco, API REST,
GraphQL e exports de ERP — devolve a MESMA estrutura ``SourceDataset``. O
resto do pipeline (mapeamento, validação, escrita, rollback) não sabe de
onde o dado veio. É isso que torna barato adicionar fonte nova depois.

A etapa 3 do assistente ("leitura automática") vive aqui: detectar
codificação, separador, planilha, colunas, tipos e quantidade de registros.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from defusedxml import ElementTree as ET

# Limites de segurança (ADR-0003 §7). Um arquivo maior que isso é recusado
# antes de qualquer parsing — não adianta validar depois de já ter alocado.
MAX_SOURCE_BYTES = 25 * 1024 * 1024
MAX_SOURCE_ROWS = 200_000

# Acima deste volume a importação não deve rodar de forma síncrona
# (ADR-0003 §8). Na Fase 1 o job é marcado e reportado; o executor em
# background entra na Fase 4.
BACKGROUND_THRESHOLD_ROWS = 100_000

FILE_SOURCE_KINDS = ('xlsx', 'csv', 'ods', 'json', 'xml', 'txt')

# Fontes declaradas no catálogo comercial mas ainda não implementadas
# (ADR-0003 §9, fases 5 e 6). Listadas para a UI poder exibi-las como
# "em breve" em vez de simplesmente não existirem.
ROADMAP_SOURCE_KINDS = (
    'sqlserver', 'mysql', 'postgresql', 'sqlite', 'oracle',
    'rest', 'graphql', 'sap', 'totvs', 'senior', 'benner', 'dominio',
)

# Sistema de origem gravado no registro importado (issue #211).
#
# É DERIVADO do `source_kind`, não pedido ao usuário: a alternativa seria um
# campo novo na API que alguém preencheria errado, ou não preencheria. Aqui,
# uma fonte nova só precisa entrar neste mapa uma vez e todo registro que ela
# produzir nasce carimbado corretamente.
#
# Distinto de `origin`: `origin` responde "veio de fora?" e tem duas respostas
# estáveis; `source_system` responde "de fora a partir de quê?" e o inventário
# cresce a cada integração. Misturar os dois obrigaria a mexer na semântica de
# `origin` toda vez que um ERP novo entrasse.
_SOURCE_SYSTEM_BY_KIND = {
    'sap': 'sap',
    'totvs': 'totvs',
    'senior': 'senior',
    'benner': 'benner',
    'dominio': 'dominio',
    'rest': 'api',
    'graphql': 'api',
    'sqlserver': 'banco',
    'mysql': 'banco',
    'postgresql': 'banco',
    'sqlite': 'banco',
    'oracle': 'banco',
}


def source_system_for(source_kind: str) -> str:
    """Sistema de origem correspondente a um ``source_kind``.

    Arquivo (xlsx/csv/ods/json/xml/txt) é sempre ``'planilha'`` — do ponto de
    vista da auditoria, o que importa é que um humano exportou e enviou um
    arquivo, não qual extensão ele tinha.

    Um `source_kind` desconhecido devolve o próprio valor normalizado em vez de
    string vazia: registro carimbado com uma fonte estranha é diagnosticável;
    registro sem carimbo nenhum é indistinguível de um bug.
    """
    kind = str(source_kind or '').strip().lower()
    if not kind:
        return ''
    if kind in FILE_SOURCE_KINDS:
        return 'planilha'
    return _SOURCE_SYSTEM_BY_KIND.get(kind, kind)


_ENCODING_CANDIDATES = ('utf-8-sig', 'utf-8', 'latin-1', 'cp1252')
_DELIMITER_CANDIDATES = (';', ',', '\t', '|')


@dataclass
class SourceDataset:
    columns: list[str]
    rows: list[dict]
    diagnostics: list[dict] = field(default_factory=list)
    detected: dict = field(default_factory=dict)
    source_lines: list[int] = field(default_factory=list)

    @property
    def total_rows(self) -> int:
        return len(self.rows)

    def signature(self) -> str:
        """Assinatura do layout de origem: mesma ordem de colunas = mesma
        assinatura. É a chave do mapeamento salvo (ADR-0003 §2.3) — o que
        faz a segunda importação do mesmo relatório legado já vir mapeada."""
        joined = '|'.join(_normalize_header(column) for column in self.columns)
        return hashlib.sha256(joined.encode('utf-8')).hexdigest()[:32]

    def sample(self, limit: int = 10) -> list[dict]:
        return self.rows[:limit]


def _normalize_header(value: object) -> str:
    text = str(value or '').strip().lower()
    for accented, plain in (
        ('á', 'a'), ('à', 'a'), ('ã', 'a'), ('â', 'a'), ('ä', 'a'),
        ('é', 'e'), ('ê', 'e'), ('è', 'e'), ('ë', 'e'),
        ('í', 'i'), ('î', 'i'), ('ì', 'i'), ('ï', 'i'),
        ('ó', 'o'), ('ô', 'o'), ('õ', 'o'), ('ò', 'o'), ('ö', 'o'),
        ('ú', 'u'), ('û', 'u'), ('ù', 'u'), ('ü', 'u'),
        ('ç', 'c'), ('ñ', 'n'),
    ):
        text = text.replace(accented, plain)
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def detect_encoding(raw: bytes) -> str:
    for candidate in _ENCODING_CANDIDATES:
        try:
            raw.decode(candidate)
            return candidate
        except (UnicodeDecodeError, LookupError):
            continue
    return 'latin-1'


def detect_delimiter(text: str) -> str:
    head = '\n'.join(text.splitlines()[:20])
    if not head:
        return ','
    try:
        return csv.Sniffer().sniff(head, delimiters=''.join(_DELIMITER_CANDIDATES)).delimiter
    except csv.Error:
        # Sniffer falha com frequência em export de ERP (aspas irregulares,
        # linhas de cabeçalho decorativas). Fallback determinístico: o
        # candidato mais frequente na primeira linha.
        first = head.splitlines()[0]
        counts = {candidate: first.count(candidate) for candidate in _DELIMITER_CANDIDATES}
        best = max(counts, key=lambda key: counts[key])
        return best if counts[best] else ','


def infer_column_types(rows: list[dict], columns: list[str]) -> dict[str, str]:
    """Tipo aparente por coluna, só para exibir no preview — a validação
    real é do descritor da entidade (catalog.py), não daqui."""
    types: dict[str, str] = {}
    for column in columns:
        values = [str(row.get(column) or '').strip() for row in rows[:200]]
        values = [value for value in values if value]
        if not values:
            types[column] = 'vazio'
            continue
        if all(re.fullmatch(r'-?\d+', value) for value in values):
            types[column] = 'inteiro'
        elif all(re.fullmatch(r'-?\d+[.,]\d+', value) for value in values):
            types[column] = 'decimal'
        elif all(re.fullmatch(r'\d{2}[/-]\d{2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2}.*', value) for value in values):
            types[column] = 'data'
        else:
            types[column] = 'texto'
    return types


class SourceError(ValueError):
    """Compatible ValueError with safe, machine-readable source diagnostics."""
    def __init__(self, message, *, row=1, field='', code='invalid_source'):
        super().__init__(f'{message} (linha {row}).')
        self.diagnostics = [{'code': code, 'row': row, 'source_row': row, 'source_line': row,
                             'field': field, 'message': str(self), 'severity': 'error',
                             'level': 'error', 'entity': ''}]


def _headers(values, *, row=1):
    columns = [str(v if v is not None else '').strip() for v in values]
    if not columns or any(not name for name in columns):
        raise SourceError('Cabeçalho vazio ou ausente', row=row, code='invalid_header')
    normalized = [_normalize_header(name) for name in columns]
    if len(set(normalized)) != len(normalized):
        raise SourceError('Cabeçalhos duplicados ou equivalentes', row=row, code='duplicate_header')
    return columns


def _finalize(columns, rows, detected, *, diagnostics=None, source_lines=None):
    columns = _headers(columns)
    if not rows:
        raise SourceError('Arquivo vazio: nenhum registro de dados')
    if len(rows) > MAX_SOURCE_ROWS:
        raise SourceError(f'Arquivo excede o limite de {MAX_SOURCE_ROWS} registros')
    detected = {**detected, 'total_rows': len(rows), 'column_count': len(columns),
                'column_types': infer_column_types(rows, columns),
                'requires_background': len(rows) >= BACKGROUND_THRESHOLD_ROWS}
    diagnostics = list(diagnostics or [])
    if detected['requires_background']:
        diagnostics.append({'level': 'warning', 'severity': 'warning', 'row': 0, 'field': '',
                            'code': 'requires_background', 'message': 'Volume exige executor assíncrono; consulte os limites operacionais.'})
    return SourceDataset(columns=columns, rows=rows, diagnostics=diagnostics, detected=detected,
                         source_lines=source_lines or list(range(1, len(rows) + 1)))


def _read_delimited(raw, *, kind):
    encoding = detect_encoding(raw)
    text = raw.decode(encoding)
    delimiter = detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
    rows, lines, diagnostics = [], [], []
    try:
        columns = _headers(next(reader, []))
        next_line = reader.line_num + 1
        for values in reader:
            line = next_line
            next_line = reader.line_num + 1
            if not values:
                diagnostics.append({'code': 'blank_source_line', 'row': line, 'source_line': line,
                                    'field': '', 'level': 'warning', 'severity': 'warning',
                                    'message': 'Linha vazia sem valores não importada.'})
                continue
            if len(values) != len(columns):
                raise SourceError('Quantidade de colunas incompatível com o cabeçalho', row=line,
                                  code='column_count_mismatch')
            rows.append(dict(zip(columns, values)))
            lines.append(line)
            if len(rows) > MAX_SOURCE_ROWS:
                raise SourceError(f'Limite de {MAX_SOURCE_ROWS} registros excedido', row=line)
    except csv.Error as exc:
        raise SourceError('CSV inválido ou corrompido', row=max(reader.line_num, 1)) from exc
    return _finalize(columns, rows, {'encoding': encoding, 'delimiter': delimiter, 'kind': kind},
                     diagnostics=diagnostics, source_lines=lines)


def _cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return '' if value is None else str(value)


def _read_spreadsheet(raw, *, kind, sheet=None):
    try:
        if kind == 'xlsx':
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            sheet_names = book.sheetnames
            if sheet and sheet not in sheet_names:
                raise SourceError('Aba solicitada não existe', field='sheet')
            chosen = sheet or sheet_names[0]
            iterator = book[chosen].iter_rows()
            header = next(iterator, ())
            values = [c.value for c in header]
            while values and values[-1] is None:
                values.pop()
            columns = _headers(values)
            rows, lines = [], []
            for number, cells in enumerate(iterator, 2):
                values = [c.value for c in cells]
                if not any(v is not None for v in values):
                    continue
                if any(getattr(c, 'data_type', '') in ('f', 'e') for c in cells):
                    raise SourceError('Fórmula ou erro de planilha exige exportação de valores', row=number)
                if any(v is not None for v in values[len(columns):]):
                    raise SourceError('Valor sem cabeçalho na planilha', row=number)
                values = values[:len(columns)] + [None] * max(0, len(columns) - len(values))
                rows.append(dict(zip(columns, map(_cell, values))))
                lines.append(number)
                if len(rows) > MAX_SOURCE_ROWS:
                    raise SourceError(f'Limite de {MAX_SOURCE_ROWS} registros excedido', row=number)
            book.close()
        else:
            import pandas as pd
            with pd.ExcelFile(io.BytesIO(raw), engine='odf') as book:
                sheet_names = book.sheet_names
                if sheet and sheet not in sheet_names:
                    raise SourceError('Aba solicitada não existe', field='sheet')
                chosen = sheet or sheet_names[0]
                frame = book.parse(chosen, header=None, dtype=str, keep_default_na=False)
                if frame.empty:
                    raise SourceError('Planilha vazia')
                columns = _headers(frame.iloc[0].tolist())
                rows = [dict(zip(columns, values)) for values in frame.iloc[1:].values.tolist()]
                lines = list(range(2, len(rows) + 2))
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError('Não foi possível ler a planilha; verifique formato e integridade') from exc
    warnings = []
    if len(sheet_names) > 1:
        warnings.append({'code': 'selected_sheet_only', 'row': 0, 'field': 'sheet',
                         'level': 'warning', 'severity': 'warning',
                         'message': 'Somente a aba selecionada será importada; confira as demais abas.'})
    return _finalize(columns, rows, {'kind': kind, 'sheet': chosen, 'sheets': sheet_names},
                     diagnostics=warnings, source_lines=lines)


class _JSONObject(list):
    pass


def _read_json(raw):
    encoding = detect_encoding(raw)
    try:
        payload = json.loads(raw.decode(encoding), object_pairs_hook=_JSONObject, parse_float=Decimal,
                             parse_constant=lambda value: (_ for _ in ()).throw(ValueError('nonfinite')))
    except (ValueError, UnicodeError) as exc:
        raise SourceError('JSON inválido') from exc
    warnings = []
    if isinstance(payload, _JSONObject):
        names = _headers([key for key, value in payload])
        wrappers = [(key, value) for key, value in payload if key in ('data', 'items', 'records', 'rows', 'results')
                    and isinstance(value, list) and not isinstance(value, _JSONObject)]
        if wrappers:
            if len(wrappers) != 1:
                raise SourceError('JSON com múltiplas listas de registros é ambíguo')
            wrapper, content = wrappers[0]
            for name in names:
                if name != wrapper:
                    warnings.append({'code': 'unmapped_envelope_field', 'row': 0, 'field': name,
                                     'level': 'warning', 'severity': 'warning',
                                     'message': 'Campo do envelope JSON não faz parte dos registros importados.'})
            payload = content
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise SourceError('JSON precisa conter uma lista de registros')
    columns, rows = [], []
    for number, item in enumerate(payload, 1):
        if not isinstance(item, _JSONObject):
            raise SourceError('JSON contém registro que não é objeto', row=number)
        names = _headers([key for key, value in item], row=number)
        row = {}
        for name, (_, value) in zip(names, item):
            if isinstance(value, (list, dict)):
                raise SourceError('Campo JSON aninhado não é suportado', row=number, field=name)
            row[name] = value
            if name not in columns:
                columns.append(name)
        rows.append(row)
    return _finalize(columns, rows, {'kind': 'json', 'encoding': encoding}, diagnostics=warnings)


def _read_xml(raw):
    try:
        root = ET.fromstring(raw)
    except Exception as exc:
        raise SourceError('XML inválido') from exc
    if (root.text or '').strip():
        raise SourceError('Texto XML fora dos registros não é suportado')
    rows, columns = [], []
    for number, record in enumerate(root, 1):
        if (record.text or '').strip() or (record.tail or '').strip():
            raise SourceError('Texto XML fora de campo não é suportado', row=number)
        pairs = list(record.attrib.items())
        for child in record:
            if list(child) or child.attrib or (child.tail or '').strip():
                raise SourceError('Campo XML aninhado ou misto não é suportado', row=number, field=str(child.tag))
            pairs.append((str(child.tag), (child.text or '').strip()))
        names = _headers([key for key, value in pairs], row=number)
        row = {name: value for name, (_, value) in zip(names, pairs)}
        rows.append(row)
        for name in names:
            if name not in columns:
                columns.append(name)
    warnings = [{'code': 'unmapped_envelope_field', 'row': 0, 'field': key, 'severity': 'warning',
                 'level': 'warning', 'message': 'Atributo da raiz XML não faz parte dos registros.'} for key in root.attrib]
    return _finalize(columns, rows, {'kind': 'xml', 'root': root.tag}, diagnostics=warnings)


def read_source(kind: str, raw: bytes, *, sheet: str | None = None) -> SourceDataset:
    """Ponto único de entrada de qualquer fonte de arquivo."""
    kind = str(kind or '').strip().lower()
    if not isinstance(raw, (bytes, bytearray)):
        raise SourceError('Conteúdo do arquivo ausente.')
    if len(raw) > MAX_SOURCE_BYTES:
        raise SourceError(
            f'Arquivo de {len(raw)} bytes excede o limite de {MAX_SOURCE_BYTES} bytes.'
        )
    if not raw:
        raise SourceError('Arquivo vazio.')
    if kind in ROADMAP_SOURCE_KINDS:
        raise ValueError(
            f'A origem "{kind}" ainda não está disponível '
            '(prevista para as fases 5/6 do Centro de Migração).'
        )
    if kind in ('csv', 'txt'):
        return _read_delimited(bytes(raw), kind=kind)
    if kind in ('xlsx', 'ods'):
        return _read_spreadsheet(bytes(raw), kind=kind, sheet=sheet)
    if kind == 'json':
        return _read_json(bytes(raw))
    if kind == 'xml':
        return _read_xml(bytes(raw))
    raise ValueError(f'Origem de dados não suportada: {kind!r}.')


def source_hash(raw: bytes) -> str:
    """Hash do arquivo — auditoria (ADR-0003 §6) e detecção de reimportação
    acidental do mesmo arquivo."""
    return hashlib.sha256(bytes(raw)).hexdigest()


def list_sources() -> list[dict]:
    return (
        [{'kind': kind, 'enabled': True, 'phase': '1'} for kind in FILE_SOURCE_KINDS]
        + [{'kind': kind, 'enabled': False, 'phase': '5/6'} for kind in ROADMAP_SOURCE_KINDS]
    )
