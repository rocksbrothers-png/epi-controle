#!/usr/bin/env python3
"""Loopback-only synthetic HTTP fixture for real Flutter runtime acceptance.

Never uses a production DSN or publishes official documents. Credentials stay in
an explicitly requested, mode-0600 local file; do not upload that file as artifact.
"""
import argparse
import base64
from contextlib import closing
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import threading
from urllib.parse import urlparse

root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    target = urlparse(os.environ.get('DATABASE_URL', ''))
    if os.environ.get('APP_ENV') != 'test' or target.scheme not in ('postgres', 'postgresql') or target.hostname not in ('localhost', '127.0.0.1') or not target.path.startswith('/portal_docs_'):
        raise SystemExit('Requires APP_ENV=test and explicitly local portal_docs_* PostgreSQL')
    from core.bootstrap import init_db
    init_db()
    import app
    from core.database import get_connection
    from http.server import ThreadingHTTPServer
    from modules.safety_library.preview import MIMES
    spec = importlib.util.spec_from_file_location('library_fixture', root/'tests/test_portal_docs_postgres.py')
    fixture = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fixture)
    with closing(get_connection()) as connection:
        tenant = fixture.seed_tenant(connection, 'SYNTHETIC LOCAL TEST NOT OFFICIAL')
        foreign = fixture.seed_tenant(connection, 'FOREIGN SYNTHETIC TEST')
    app._set_bootstrap_state(ready=True, error_code='', error_message='')
    server = ThreadingHTTPServer(('127.0.0.1', 0), app.EpiHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = 'http://127.0.0.1:'+str(server.server_port)
    versions = []
    from docx import Document
    from openpyxl import Workbook
    import xlwt
    for extension in ('.docx', '.xlsx', '.xls', '.pdf', '.png'):
        content = io.BytesIO()
        if extension == '.docx':
            document = Document()
            for page in range(5):
                if page:
                    document.add_page_break()
                document.add_heading('Synthetic safety page '+str(page+1), 0)
                document.add_paragraph('Read authorized safety instructions.')
            document.save(content)
        elif extension == '.xlsx':
            workbook = Workbook()
            workbook.active.title = 'POR'
            workbook.create_sheet('ENG')
            for sheet in workbook:
                sheet.merge_cells('A1:C1')
                sheet['A1'] = 'Distinct sheet '+sheet.title
                sheet['A2'] = 'Safety guidance'
                sheet.print_area = 'A1:D12'
            workbook.save(content)
        elif extension == '.xls':
            workbook = xlwt.Workbook()
            sheet = workbook.add_sheet('Glove matrix')
            for row in range(100):
                sheet.write(row, 0, 'Distinct safety row '+str(row))
                sheet.write(row, 1, 'Wear appropriate gloves')
            workbook.save(content)
        else:
            content.write(fixture.pdf() if extension == '.pdf' else fixture.png())
        data = fixture.upload_data(tenant, title='Fixture isolada '+extension[1:].upper(), filename='synthetic'+extension, mime=MIMES[extension], content_base64=base64.b64encode(content.getvalue()).decode())
        status, result, _ = fixture.request(base, 'upload', data, tenant=tenant)
        if status != 200:
            raise RuntimeError('Synthetic fixture upload failed: '+str(result))
        version = result['version']
        status, result, _ = fixture.request(base, 'publish', {'document_id':version['document_id'], 'version_id':version['id'], 'justification':'Synthetic isolated test, never official content'}, tenant=tenant)
        if status != 200:
            raise RuntimeError('Synthetic fixture publication failed: '+str(result))
        versions.append(version)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with open(output.with_suffix('.tmp'), 'w', opener=lambda path, flags: os.open(path, flags, 0o600)) as handle:
        json.dump({'pid':os.getpid(), 'base_url':base, 'tenant':tenant, 'foreign':foreign, 'versions':versions}, handle)
    output.with_suffix('.tmp').replace(output)
    output.chmod(0o600)
    print('Synthetic local HTTP fixture ready', flush=True)
    threading.Event().wait()


if __name__ == '__main__':
    main()
